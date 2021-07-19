import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple  # todo rewrite so I only use one package
from xlsxwriter.utility import xl_rowcol_to_cell


def get_data_from_excel_table_by_colname(table, workbook, col_name):
    start_range, end_range = table.ref.split(':')
    start_range = coordinate_to_tuple(start_range)
    end_range = coordinate_to_tuple(end_range)

    id_ = table.column_names.index(col_name)
    data_range = (start_range[1] + id_, start_range[0] + 1, start_range[1] + id_, end_range[0])
    data_range_ = xl_rowcol_to_cell(data_range[1] - 1, data_range[0] - 1) + ':' + xl_rowcol_to_cell(data_range[3] - 1,
                                                                                                    data_range[2] - 1)
    data_tuple = workbook.active[data_range_]

    return [data_tuple[i][0].value for i in range(len(data_tuple))]


def generate_weights_json(excel_file_name, strat_names, output_file_name):
    wb = load_workbook(filename=excel_file_name)
    pos_list = wb.worksheets[0].tables['Position_list']
    col_names = ['Financial Instrument'] + [strat + ' %' for strat in strat_names]

    df_ = pd.concat([pd.Series(data=get_data_from_excel_table_by_colname(pos_list, wb, col_name), name=col_name)
                     for col_name in col_names], axis=1)
    df_.to_json(output_file_name)


if __name__ == '__main__':
    input_file_name = 'C:/Users/Anton/PycharmProjects/Portfolio_rebalancing_tool/outputs/E18072021_clean_filled.xlsx'
    strats_list1 = ['Redhead', 'Workhorse', 'Safe'] # todo - pass this programmatically
    strats_list2 = ['Crypto', 'MJ', 'Materials and Defensives', 'Researched Growth Companies', 'Momentum Hype']
    output_file_name = 'C:/Users/Anton/PycharmProjects/Portfolio_rebalancing_tool/jsons/save_weights_E18072021.json'
    generate_weights_json(input_file_name, strats_list2, output_file_name)
